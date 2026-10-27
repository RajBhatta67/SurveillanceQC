"""Reproducible LOCAL audit of the supplied handoff, not a public data fixture.

Run from an installed SurveillanceQC environment:
python tools/check_research_handoff.py --handoff PATH --output NEW_DIRECTORY
Optional workbook comparison requires openpyxl; no archived analysis is executed.
"""
import argparse
from collections import Counter, defaultdict
import csv
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import platform
import re
from typing import NamedTuple

from surveillanceqc import Config, check_csv, __version__

DISEASES = ('measles', 'rubella')
COMPONENTS = ('clinical', 'epi_linked', 'lab_confirmed', 'total')
MONTH_COUNTS = ('measles_suspect', *(f'{d}_{c}' for d in DISEASES for c in COMPONENTS), 'discarded')
YEAR_COUNTS = ('total_suspected_measles_rubella_cases', *(f'{d}_{c}' for d in DISEASES for c in COMPONENTS), 'discarded_cases')
SOURCES = {
    'original_monthly': ('baseline/data/measles_rubella_monthly.csv', 'month', 'original supplied FTI research snapshot'),
    'original_annual': ('baseline/data/measles_rubella_yearly.csv', 'year', 'original supplied FTI research snapshot'),
    'original_merged': ('baseline/data/measles_rubella_merged.csv', 'month', 'supplied upstream merged product, not raw WHO'),
    'reference_monthly': ('source_audit/cases_month.csv', 'month', 'frozen TidyTuesday reference, not primary input'),
    'reference_annual': ('source_audit/cases_year.csv', 'year', 'frozen TidyTuesday reference, not primary input'),
    'later_monthly': ('validation2026/data/who_monthly_2026.csv', 'month', 'archived derived September 2026 conversion, original converter absent'),
}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream)
        rows = list(reader)
        if any(None in row or None in row.values() for row in rows):
            raise ValueError(f'Malformed record in {path}')
        return reader.fieldnames, rows


def write_csv(path, columns, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def write_json(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=False, allow_nan=False) + '\n', encoding='utf-8')


class NumericValue(NamedTuple):
    status: str  # valid, missing, invalid: invalid must never become missing
    value: Decimal | None


def parse_number(value, *, whole=False):
    text = "" if value is None else str(value).strip()
    if text in ("", "NA"):
        return NumericValue("missing", None)
    try:
        # Same explicit syntax as the core checker; do not accept Decimal's
        # extra syntax (e.g. underscores), or nonfinite tokens.
        if not re.fullmatch(r"[+-]?[0-9]+(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?", text):
            raise InvalidOperation
        value = Decimal(text)
        if not value.is_finite() or (whole and value != value.to_integral_value()):
            raise InvalidOperation
    except InvalidOperation:
        return NumericValue("invalid", None)
    return NumericValue("valid", value)


def number(value, *, whole=False):
    """Convenience for guarded arithmetic; use status for missingness reporting."""
    return parse_number(value, whole=whole).value


def numeric_issues(rows):
    issues = []
    for record, row in enumerate(rows, start=2):
        for column, token in row.items():
            if column in ('iso3', 'country', 'region'):
                continue
            whole = column in (*MONTH_COUNTS, *YEAR_COUNTS, 'year', 'month')
            if parse_number(token, whole=whole).status == 'invalid':
                issues.append(dict(source_record=record, column=column, value=token,
                                   reason='invalid whole number' if whole else 'invalid finite number'))
    return issues


def period_text(row, frequency):
    """Join explicit calendar fields; never drop/repair an invalid source period."""
    y = number(row['year'])
    m = number(row['month']) if frequency == 'month' else None
    if y is None or y != y.to_integral_value() or not 1 <= y <= 9999:
        return f"INVALID_YEAR:{row['year']}"
    if frequency == 'year':
        return f'{int(y):04d}'
    if m is None or m != m.to_integral_value() or not 1 <= m <= 12:
        return f"INVALID_MONTH:{row['year']}:{row['month']}"
    return f'{int(y):04d}-{int(m):02d}'


def normalize(columns, rows, frequency):
    if 'qc_period' in columns:
        raise ValueError('qc_period already present')
    return [dict(row, qc_period=period_text(row, frequency)) for row in rows]


def eligible(row, disease):
    values = [number(row[f'{disease}_{c}'], whole=True) for c in COMPONENTS]
    return all(v is not None and v >= 0 and v == v.to_integral_value() for v in values)


def key(row, frequency):
    # A malformed key is retained as evidence but cannot align comparisons.
    if not row['iso3'].strip() or row['iso3'].strip() == 'NA':
        return None
    if period_text(row, frequency).startswith('INVALID_'):
        return None
    year = str(int(number(row['year'])))
    return (row['iso3'], year, str(int(number(row['month'])))) if frequency == 'month' else (row['iso3'], year)


def index_rows(rows, frequency):
    indexed, skipped = {}, []
    for record, row in enumerate(rows, start=2):
        identity = key(row, frequency)
        if identity is None:
            skipped.append(dict(source_record=record, reason='invalid or missing calendar/identity key'))
        elif identity in indexed:
            raise ValueError('Comparison requires unique keys; cannot silently overwrite duplicates')
        else:
            indexed[identity] = row
    return indexed, skipped


def compare_rows(left, right, frequency):
    a, skipped_left = index_rows(left, frequency)
    b, skipped_right = index_rows(right, frequency)
    shared = sorted(a.keys() & b.keys())
    differences, skipped_numbers = Counter(), []
    examples = []
    # Empty input is a supported comparison, not an index error.
    fields = sorted((set(left[0]) if left else set()) &
                    (set(right[0]) if right else set()) - {'qc_period'})
    for identity in shared:
        for field in fields:
            v, w = a[identity][field], b[identity][field]
            if field not in ('region', 'country', 'iso3'):
                whole = field in (*MONTH_COUNTS, *YEAR_COUNTS, 'year', 'month')
                x, y = parse_number(v, whole=whole), parse_number(w, whole=whole)
                if 'invalid' in (x.status, y.status):
                    skipped_numbers.append(dict(key=identity, column=field, left=v, right=w,
                                                reason='invalid numeric operand'))
                    continue
                # Missing/missing equivalence is explicit, never missing == 0.
                same = x.status == y.status and x.value == y.value
            else:
                same = v == w
            if not same:
                differences[field] += 1
                if len(examples) < 20:
                    examples.append(dict(key=identity, column=field, left=v, right=w))
    return dict(shared_keys=len(shared), left_only_keys=sorted(a.keys()-b.keys()),
                right_only_keys=sorted(b.keys()-a.keys()), differences_by_column=dict(differences),
                first_20_differences=examples, skipped_left_records=skipped_left,
                skipped_right_records=skipped_right, skipped_numeric_comparisons=skipped_numbers,
                analysis_complete=not (skipped_left or skipped_right or skipped_numbers))


def reconcile_annual(yearly, monthly):
    annual_index, skipped_annual = index_rows(yearly, 'year')
    monthly_index, skipped_monthly = index_rows(monthly, 'month')
    yearly = [r for k, r in annual_index.items() if 2012 <= int(k[1]) <= 2024]
    grouped = defaultdict(list)
    for k, row in monthly_index.items():
        if 2012 <= int(k[1]) <= 2024:
            grouped[k[:2]].append(row)
    summary = dict(analysis_complete=not (skipped_annual or skipped_monthly),
                   skipped_annual_records=skipped_annual, skipped_monthly_records=skipped_monthly)
    for disease in DISEASES:
        fields = {}
        for category in COMPONENTS:
            field = f'{disease}_{category}'
            no_monthly = zeros_no_monthly = mismatches = observed = 0
            skipped_comparisons = []
            for row in yearly:
                identity = key(row, 'year')
                cells = [parse_number(m[field], whole=True) for m in grouped[identity]]
                annual = parse_number(row[field], whole=True)
                # Do not sum only valid entries if another entry is invalid:
                # that would fabricate a partial total and a false mismatch.
                if annual.status != 'valid' or any(c.status == 'invalid' for c in cells):
                    skipped_comparisons.append(dict(key=identity, field=field,
                        reason='annual operand missing/invalid or monthly operand invalid'))
                    continue
                values = [c.value for c in cells if c.status == 'valid']
                if not values:
                    no_monthly += 1
                    zeros_no_monthly += annual.value == 0
                else:
                    observed += 1
                    # Retain negative source values in this descriptive sum.
                    mismatches += annual.value != sum(values, Decimal(0))
            fields[field] = dict(observed_monthly_sums=observed, unequal_available_monthly_sums=mismatches,
                no_monthly_values=no_monthly, annual_zero_without_monthly_values=zeros_no_monthly,
                skipped_comparisons=skipped_comparisons, analysis_complete=not skipped_comparisons)
            if skipped_comparisons:
                summary['analysis_complete'] = False
        summary[disease] = fields
    suspected = 'total_suspected_measles_rubella_cases'
    discarded = 'discarded_cases'
    rate = 'discarded_non_measles_rubella_cases_per_100000_total_population'
    annual_cells = [(parse_number(r[suspected], whole=True), parse_number(r[discarded], whole=True),
                     parse_number(r[rate])) for r in yearly]
    invalid = sum(c.status == 'invalid' for row in annual_cells for c in row)
    if invalid:
        summary['analysis_complete'] = False
    summary['annual_2012_2024'] = dict(rows=len(yearly),
        suspected_missing=sum(a.status == 'missing' for a,b,c in annual_cells),
        discarded_missing=sum(b.status == 'missing' for a,b,c in annual_cells),
        discarded_negative=sum(b.status == 'valid' and b.value < 0 for a,b,c in annual_cells),
        discarded_present_rate_missing=sum(b.status == 'valid' and c.status == 'missing' for a,b,c in annual_cells),
        invalid_numeric_cells=invalid, analysis_complete=invalid == 0)
    return summary


# Reviewed workbook labels map explicitly to semantic fields. CSV column order
# is never consulted. New workbook schemas require a new reviewed mapping.
WORKBOOK_FIELDS = {
    'Region': 'region', 'Country': 'country', 'ISO3': 'iso3',
    'Year': 'year', 'Month': 'month',
    'Measles \nsuspect': 'measles_suspect',
    'Measles \nclinical': 'measles_clinical',
    'Measles \nepi-linked': 'measles_epi_linked',
    'Measles \nlab-confirmed': 'measles_lab_confirmed',
    'Measles \ntotal': 'measles_total',
    'Rubella \nclinical': 'rubella_clinical',
    'Rubella \nepi-linked': 'rubella_epi_linked',
    'Rubella \nlab-confirmed': 'rubella_lab_confirmed',
    'Rubella\nTotal': 'rubella_total', 'Discarded': 'discarded',
}


def extract_workbook(workbook):
    ws = workbook['WEB']
    records = ws.iter_rows()
    header = next(records)
    labels = [c.value for c in header if c.value is not None]
    if len(labels) != len(set(labels)) or set(labels) != set(WORKBOOK_FIELDS):
        raise ValueError('Workbook header differs from reviewed mapping')
    positions = {c.value: i for i,c in enumerate(header) if c.value is not None}
    unused = set(range(len(header))) - set(positions.values())
    extracted, source_rows = [], []
    for record, cells in enumerate(records, start=2):
        if all(c.value is None for c in cells):
            continue
        if any(cells[i].value is not None for i in unused):
            raise ValueError('Unexpected nonempty unheaded workbook column')
        if any(c.data_type == 'f' for c in cells):
            raise ValueError('Formula encountered; no silent use of cached values')
        extracted.append({field: '' if cells[positions[label]].value is None else str(cells[positions[label]].value)
                          for label, field in WORKBOOK_FIELDS.items()})
        source_rows.append(record)
    return extracted, source_rows


def workbook_comparison(root, output, derived):
    from openpyxl import load_workbook
    import openpyxl
    path = root / 'research_v6/validation2026/data/b5efa946-404-table-web-epi-curve-data.xlsx'
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        extracted, source_excel_rows = extract_workbook(wb)
        columns = tuple(WORKBOOK_FIELDS.values())
        converted = output / 'normalized/later_workbook_new_extraction.csv'
        write_csv(converted, columns, extracted)
        write_json(output/'evidence/later_workbook_row_map.json', source_excel_rows)
        result = compare_rows(extracted, derived, 'month')
        result.update(source_sha256=sha(path), new_extraction_sha256=sha(converted),
                      original_converter_recovered=False, extractor='new handoff audit, not original research code',
                      sheet='WEB', workbook_rows=len(extracted), archived_csv_rows=len(derived),
                      version_label=wb['Read Me']['B12'].value,
                      definition_cell='Read Me!B10', openpyxl_version=openpyxl.__version__,
                      numeric_equivalence_not_byte_identity=True)
        return result
    finally:
        wb.close()


def run(handoff, output, verify_workbook=False):
    root, output = handoff.resolve(), output.resolve()
    if output == root or output.is_relative_to(root):
        raise ValueError('Output must be outside preserved handoff')
    if output.exists() and any(output.iterdir()):
        raise ValueError('Output directory must be new or empty')
    output.mkdir(parents=True, exist_ok=True)
    before = {p.relative_to(root).as_posix(): sha(p) for p in root.rglob('*') if p.is_file()}
    verified = []
    for line in (root/'SHA256SUMS.txt').read_text().splitlines():
        expected, name = line.split('  ', 1)
        if before.get(name) != expected:
            raise ValueError(f'Handoff hash mismatch: {name}')
        verified.append(name)
    manifest = json.loads((root/'FILE_MANIFEST.json').read_text())
    for item in manifest['files']:
        if before.get(item['path']) != item['sha256'] or (root/item['path']).stat().st_size != item['bytes']:
            raise ValueError('Manifest mismatch: ' + item['path'])
    report = dict(tool_version=__version__, python_version=platform.python_version(),
                  audit_script_sha256=sha(Path(__file__)),
                  evidence='fresh checks on supplied files, not a rerun of research estimates',
                  sha256_entries_verified=len(verified), manifest_entries_verified=len(manifest['files']),
                  sources={}, conversions=[], comparisons={}, originals_unchanged=None)
    loaded = {}
    negative_evidence = []
    invalid_evidence = []
    for label, (relative, frequency, provenance) in SOURCES.items():
        source = root/'research_v6'/relative
        columns, rows = read_csv(source)
        invalid = numeric_issues(rows)
        invalid_evidence.extend(dict(item, source=label) for item in invalid)
        normalized = normalize(columns, rows, frequency)
        destination = output/'normalized'/f'{label}.csv'
        write_csv(destination, [*columns, 'qc_period'], normalized)
        counts = MONTH_COUNTS if frequency == 'month' else YEAR_COUNTS
        cfg = Config(country='iso3', period='qc_period', counts=counts,
                     frequency=frequency, missing_tokens=('NA',))
        core = check_csv(destination, cfg)
        write_json(output/'reports'/f'{label}.json', core)
        codes = Counter(f['code'] for f in core['findings'])
        info = dict(path=source.relative_to(root).as_posix(), provenance=provenance,
                    source_sha256=sha(source), rows=len(rows), years=sorted({r['year'] for r in rows}),
                    config=core['config'], findings_by_code=dict(codes), summary=core['summary'],
                    analysis_complete=core['analysis_complete'],
                    invalid_numeric_cells=len(invalid), audit_analysis_complete=not invalid and core['analysis_complete'], subset={})
        report['conversions'].append(dict(source=info['path'], output=destination.relative_to(output).as_posix(),
            source_sha256=sha(source), output_sha256=sha(destination),
            rule='Retain every source field, row order, and count token; append qc_period by joining integer calendar year/month (YYYY-MM) or year (YYYY). No filtering, imputation, or count changes. Blank/NA tokens interpreted only by explicitly configured checks.',
            row_map='Output CSV record ordinal equals source CSV record ordinal; header is record 1.'))
        for row_no, row in enumerate(rows, start=2):
            for col in counts:
                n = number(row[col], whole=True)
                if n is not None and n < 0:
                    negative_evidence.append(dict(source=label, source_record=row_no, iso3=row['iso3'],
                        country=row['country'], year=row['year'], month=row.get('month', ''), column=col, value=row[col]))
        for disease in DISEASES:
            # Check subset ONLY on complete, integral, nonnegative four-field records.
            keep = [(i, row) for i, row in enumerate(normalized, start=2)
                    if eligible(row, disease) and key(row, frequency) is not None]
            invalid_operand_records = [i for i,row in enumerate(normalized, start=2)
                if any(parse_number(row[f'{disease}_{c}'], whole=True).status == 'invalid' for c in COMPONENTS)
                or key(row, frequency) is None]
            subset_file = output/'normalized'/f'{label}_{disease}_eligible.csv'
            write_csv(subset_file, [*columns, 'qc_period'], [row for _, row in keep])
            mapping = {i: source_no for i, (source_no, _) in enumerate(keep, start=2)}
            basis = ('Same disease/ISO3/calendar period/release; clinical + epi_linked + lab_confirmed = total. '
                     'Evidence: supplied September WHO monthly workbook Read Me!B10 and pinned TidyTuesday dictionary; '
                     'historical June workbook absent. Only complete nonnegative four-field rows eligible.')
            subset_cfg = Config(country='iso3', period='qc_period', frequency=frequency,
                counts=tuple(f'{disease}_{c}' for c in COMPONENTS), missing_tokens=('NA',),
                lab=f'{disease}_lab_confirmed', total=f'{disease}_total', subset_basis=basis)
            subset_core = check_csv(subset_file, subset_cfg)
            subset_findings = []
            for f in subset_core['findings']:
                if f['code'] == 'SUBSET_EXCEEDS_TOTAL':
                    subset_findings.append(dict(f, projection_rows=f['rows'], rows=[mapping[i] for i in f['rows']]))
            sum_mismatches = []
            for source_no, row in keep:
                parts = [number(row[f'{disease}_{c}']) for c in COMPONENTS]
                if sum(parts[:3], Decimal(0)) != parts[3]:
                    sum_mismatches.append(dict(source_record=source_no, iso3=row['iso3'], period=row['qc_period']))
            result = dict(eligible_rows=len(keep), skipped_rows=len(rows)-len(keep),
                basis=basis, core_config=subset_core['config'],
                analysis_complete=not invalid_operand_records, coverage_complete=len(keep) == len(rows),
                skipped_invalid_records=invalid_operand_records,
                subset_violation_count=len(subset_findings), subset_findings=subset_findings,
                separate_sum_identity_audit_mismatches=sum_mismatches,
                scope='Only subset findings from the filtered core run are interpreted. Gap warnings on filtered projections are intentionally ignored. Sum identity is a separate audit, not a built-in v0.1 check.')
            write_json(output/'reports'/f'{label}_{disease}_subset.json', result)
            write_json(output/'evidence'/f'{label}_{disease}_eligible_source_records.json', list(mapping.values()))
            info['subset'][disease] = {k: v for k, v in result.items() if k in ('eligible_rows','skipped_rows','subset_violation_count','analysis_complete','coverage_complete','skipped_invalid_records')}
            info['subset'][disease]['sum_identity_mismatches'] = len(sum_mismatches)
        report['sources'][label] = info
        loaded[label] = rows
    write_csv(output/'evidence/negative_counts.csv',
              ['source','source_record','iso3','country','year','month','column','value'], negative_evidence)
    write_json(output/'evidence/invalid_numeric_values.json', invalid_evidence)
    report['comparisons']['monthly_reference'] = compare_rows(loaded['original_monthly'],loaded['reference_monthly'],'month')
    report['comparisons']['annual_reference'] = compare_rows(loaded['original_annual'],loaded['reference_annual'],'year')
    report['comparisons']['merged_monthly_fields'] = compare_rows(loaded['original_monthly'],loaded['original_merged'],'month')
    report['comparisons']['original_vs_later_monthly'] = compare_rows(loaded['original_monthly'],loaded['later_monthly'],'month')
    report['annual_monthly_2012_2024'] = reconcile_annual(loaded['original_annual'],loaded['original_monthly'])
    report['investigation'] = {}
    for label in SOURCES:
        negatives = [r for r in negative_evidence if r['source'] == label]
        report['investigation'][label] = dict(negative_cells=len(negatives),
            distinct_negative_records=len({r['source_record'] for r in negatives}),
            countries_with_negatives=sorted({r['iso3'] for r in negatives}))
    residual_rows = []
    for i, row in enumerate(loaded['original_annual'], start=2):
        discarded = number(row['discarded_cases'], whole=True)
        if discarded is not None and discarded < 0:
            values = [number(row[c], whole=True) for c in ('total_suspected_measles_rubella_cases','measles_total','rubella_total')]
            residual = values[0] - values[1] - values[2] if all(v is not None for v in values) else None
            residual_rows.append(dict(source_record=i, iso3=row['iso3'], year=row['year'],
                discarded=row['discarded_cases'], suspected_minus_both_disease_totals=str(residual),
                matches_residual=discarded == residual if residual is not None else None,
                analysis_complete=residual is not None))
    write_json(output/'evidence/annual_negative_residuals.json', residual_rows)
    report['investigation']['annual_negative_residual_matches'] = sum(r['matches_residual'] is True for r in residual_rows)
    if verify_workbook:
        report['comparisons']['later_workbook'] = workbook_comparison(root,output,loaded['later_monthly'])
    report['analysis_complete'] = (
        all(info['audit_analysis_complete'] for info in report['sources'].values())
        and all(c['analysis_complete'] for c in report['comparisons'].values())
        and report['annual_monthly_2012_2024']['analysis_complete']
        and all(r['analysis_complete'] for r in residual_rows))
    report['invalid_numeric_cells'] = len(invalid_evidence)
    report['workbook_check_requested'] = verify_workbook
    after = {p.relative_to(root).as_posix(): sha(p) for p in root.rglob('*') if p.is_file()}
    report['originals_unchanged'] = before == after
    report['preserved_file_count'] = len(before)
    if not report['originals_unchanged']:
        raise ValueError('Preserved handoff changed during audit')
    write_json(output/'evidence/preserved_files_sha256.json', before)
    write_json(output/'summary.json', report)
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--handoff', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    parser.add_argument('--verify-workbook', action='store_true')
    args = parser.parse_args()
    report = run(args.handoff, args.output, args.verify_workbook)
    print(json.dumps({name: info['findings_by_code'] for name,info in report['sources'].items()}, indent=2))
    print('Originals unchanged:', report['originals_unchanged'])
    print('Audit analysis complete:', report['analysis_complete'])
    print('Invalid numeric cells:', report['invalid_numeric_cells'])
    return int(not report['analysis_complete'])


if __name__ == '__main__':
    raise SystemExit(main())
